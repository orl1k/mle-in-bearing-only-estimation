import pandas as pd
import numpy as np
from collections import namedtuple


def bd_func(data, params):
    b, d, c, v = params
    n = len(data[0])
    vx, vy = convert_from_polar(v, c)
    r_y = 1000 * d * np.sin(b) + vy * np.array(range(n)) - data[1]
    r_x = 1000 * d * np.cos(b) + vx * np.array(range(n)) - data[0]
    angle = np.arctan2(r_y, r_x)
    return angle


def bd_func_jac(data, params):
    J = []
    b, d, c, v = params
    n = len(data[0])

    vx, vy = convert_from_polar(v, c)
    x, y = convert_from_polar(d, b)

    r_y = 1000 * y + vy * np.arange(n) - data[1]
    r_x = 1000 * x + vx * np.arange(n) - data[0]
    R2 = r_x ** 2 + r_y ** 2

    J.append((1000 * (x * r_x + y * r_y)) / R2)
    J.append((1000 * (np.sin(b) * r_x - np.cos(b) * x * r_y)) / R2)
    J.append((vx * r_x + vy * r_y) * np.arange(n) / R2)
    J.append((np.sin(c) * r_x - np.cos(c) * r_y) * np.arange(n) / R2)
    return np.array(J).T


def xy_func(data, params):
    x, y, vx, vy = params
    r_y = 1000 * y + vy * data[2] - data[1]
    r_x = 1000 * x + vx * data[2] - data[0]
    b = np.arctan2(r_y, r_x)
    return b


def xy_func_jac(data, params):
    J = []
    x, y, vx, vy = params
    r_y = 1000 * y + vy * data[2] - data[1]
    r_x = 1000 * x + vx * data[2] - data[0]
    R2 = r_x ** 2 + r_y ** 2

    J.append(-1000 * r_y / R2)
    J.append(1000 * r_x / R2)
    J.append(-(data[2] * r_y) / R2)
    J.append((data[2] * r_x) / R2)
    return np.array(J).T


def dist_func(r_obs, r_tar):
    return np.linalg.norm(r_obs - r_tar, 2, axis=0) / 1000


def get_random_p0(seed=None):
    rng = np.random.RandomState(seed)
    b = 0
    d = rng.uniform(5, 50)
    c = rng.uniform(0, 360)
    v = rng.uniform(5, 25)
    return [b, d, c, v]


def to_bearing(a):
    """ угол в радианах -> пеленг в радианах """
    a = (np.pi / 2) - ((2 * np.pi + a) * (a < 0) + a * (a >= 0))
    return a if a >= 0 else a + 2 * np.pi


def to_angle(b):
    """ пеленг в радианах -> угол в радианах """
    angle = (np.pi / 2) - b
    return angle if abs(angle) <= np.pi else (2 * np.pi) + angle


def convert_to_polar(x, y):
    a = np.linalg.norm([x, y], 2, axis=0)
    b = np.arctan2(y, x)
    return [b, a]


def convert_from_polar(r, a):
    x = r * np.cos(a)
    y = r * np.sin(a)
    return [x, y]


def convert_to_bdcv(params):
    x, y, vx, vy = params
    b, d = convert_to_polar(x, y)
    c, v = convert_to_polar(vx, vy)
    b = np.degrees(to_bearing(b))
    c = np.degrees(to_bearing(c))
    return np.array([b, d, c, v])


def convert_to_xy(params):
    b, d, c, v = params
    b = to_angle(np.radians(b))
    c = to_angle(np.radians(c))
    x, y = convert_from_polar(d, b)
    vx, vy = convert_from_polar(v, c)
    return np.array((x, y, vx, vy))


def get_random_p0(seed=None):
    rng = np.random.RandomState(seed)
    b = 0
    d = rng.uniform(5, 50)
    c = rng.uniform(0, 180)
    v = rng.uniform(5, 15)
    return [b, d, c, v]

def df_to_docx(df, path):

    import docx

    doc = docx.Document()
    t = doc.add_table(df.shape[0] + 2, df.shape[1])
    t.style = "Table Grid"

    for j in range(df.shape[-1]):
        t.cell(0, j).text = df.columns[j]

    for i in range(df.shape[0]):
        for j in range(df.shape[-1]):
            t.cell(i + 1, j).text = str(df.values[i, j])

    doc.save(path)